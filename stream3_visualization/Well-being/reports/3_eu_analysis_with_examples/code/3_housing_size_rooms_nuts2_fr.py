"""
3_housing_size_rooms_nuts2_fr.py — Household Size × Rooms by NUTS2 + Income Decile (France)
=============================================================================================
Computes weighted probability matrices of household size (HB120) vs. number of rooms
(HH030) disaggregated by:
  - NATIONAL level  (all France pooled)
  - NUTS2 region    (DB040)
  - income decile   (based on equivalized disposable income, HY020 / OECD scale)

Two probability types per matrix cell:
  - joint:       P(hh_size = h, rooms = r | level, decile)
  - conditional: P(rooms = r | hh_size = h, level, decile)

Income equivalization: OECD modified scale (1.0 / 0.5 / 0.3).
Decile thresholds computed at national level per year using DB090 weights.
Years are pooled; decile thresholds are re-computed independently per year.

Outputs  →  outputs/graphs/EU-SILC/housing_size_rooms_nuts2/
  - Long CSV                    : all levels, all deciles, joint + conditional
  - fr_national_matrices.xlsx   : 1 sheet per decile ('ALL' + D01..D10)
  - fr_nuts2_matrices.xlsx      : 1 sheet per (region × decile)
  - heatmaps/                   : PNG per (level, decile)

Data Source: EU-SILC Cross-sectional 2004-2023
Files used per year:
  H file : HB010 HB020 HB030 HB120 HH030 HY020
  R file : RB010 RB020 RB030 RB081 RB082   (age for OECD equiv scale)
  D file : DB010 DB020 DB030 DB040 DB090   (NUTS2, weight)
"""

import os
import warnings

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd

warnings.filterwarnings('ignore')

# ============================================================================
# CONFIGURATION
# ============================================================================

BASE_DATA_PATH = (
    r"C:/Users/valentin.stuhlfauth/OneDrive - univ-lyon2.fr"
    r"/1_WSL/1_EWBI/0_data/EU-SILC"
    r"/_Cross_2004-2023_full_set/_Cross_2004-2023_full_set"
)

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR    = os.path.abspath(os.path.join(CURRENT_DIR, '..'))
OUTPUT_BASE = os.path.join(
    BASE_DIR, 'outputs', 'graphs', 'EU-SILC', 'housing_size_rooms_nuts2'
)
os.makedirs(OUTPUT_BASE, exist_ok=True)

COUNTRY   = 'FR'
YEARS     = [2021, 2022, 2023]   # pooled for statistical robustness

# Values above these caps are grouped into a "N+" bucket
MAX_HH_SIZE = 5    # household sizes > 5 → labelled "5+"
MAX_ROOMS   = 6    # EU-SILC HH030 is already capped at 6 in source data → labelled "6+"
N_DECILES   = 10

NATIONAL_LABEL = 'FR_national'


# ============================================================================
# FILE PATH HELPERS
# ============================================================================

def file_code(country: str, year: int) -> str:
    if country == 'EL' and year <= 2007:
        return 'GR'
    return country


def _path(letter: str, country: str, year: int) -> str:
    fc = file_code(country, year)
    return f"{BASE_DATA_PATH}/{country}/{year}/UDB_c{fc}{str(year)[-2:]}{letter}.csv"


# ============================================================================
# OECD MODIFIED EQUIVALENCE SCALE HELPERS
# ============================================================================

def _oecd_person_weight(age) -> float:
    """Weight for one household member (reference person override applied later)."""
    if pd.isna(age):
        return 0.5
    try:
        return 0.3 if int(age) < 14 else 0.5
    except Exception:
        return 0.5


def compute_oecd_equiv_size(year: int, country: str) -> pd.DataFrame | None:
    """
    Read R file and compute the OECD modified equivalence scale per household.
    Returns DataFrame with columns: HB010 HB020 HB030 equiv_size
    """
    path = _path('R', country, year)
    if not os.path.exists(path):
        return None

    age_col_sets = [
        ['RB010', 'RB020', 'RB030', 'RB081', 'RB082'],
        ['RB010', 'RB020', 'RB030', 'RB082'],
        ['RB010', 'RB020', 'RB030', 'RB081'],
        ['RB010', 'RB020', 'RB030'],
    ]
    df = None
    for cols in age_col_sets:
        try:
            df = pd.read_csv(path, usecols=cols, on_bad_lines='skip')
            break
        except (ValueError, KeyError):
            continue
    if df is None:
        return None

    if 'RB081' in df.columns and 'RB082' in df.columns:
        df['age'] = df['RB081'].fillna(df['RB082'])
    elif 'RB081' in df.columns:
        df['age'] = df['RB081']
    elif 'RB082' in df.columns:
        df['age'] = df['RB082']
    else:
        df['age'] = np.nan

    # Derive household key from personal key (last 2 digits = person suffix)
    df['RB030'] = df['RB030'].fillna(0).astype(str)
    df['hh_key'] = df['RB030'].str[:-2]

    df['person_w'] = df['age'].apply(_oecd_person_weight)

    # Sort by age descending so oldest (reference person) comes first
    df = df.sort_values(
        by=['RB010', 'RB020', 'hh_key', 'age'],
        ascending=[True, True, True, False],
    )
    df['rank'] = df.groupby(['RB010', 'RB020', 'hh_key']).cumcount()
    df.loc[df['rank'] == 0, 'person_w'] = 1.0  # reference person always 1.0

    equiv = (
        df.groupby(['RB010', 'RB020', 'hh_key'])['person_w']
        .sum()
        .reset_index()
        .rename(columns={'RB010': 'HB010', 'RB020': 'HB020',
                         'hh_key': 'HB030', 'person_w': 'equiv_size'})
    )
    equiv['HB010'] = equiv['HB010'].astype(str)
    equiv['HB020'] = equiv['HB020'].astype(str)
    equiv['HB030'] = equiv['HB030'].astype(str)
    return equiv


# ============================================================================
# DATA LOADING
# ============================================================================

def load_household_full(year: int, country: str) -> pd.DataFrame | None:
    """
    H file: HB010 HB020 HB030 HB120 HH030 HY020
    HB120 = household size, HH030 = number of rooms, HY020 = disposable income
    """
    path = _path('H', country, year)
    if not os.path.exists(path):
        print(f"  [WARN] H file not found: {path}")
        return None
    try:
        df = pd.read_csv(
            path,
            usecols=['HB010', 'HB020', 'HB030', 'HB120', 'HH030', 'HY020'],
            on_bad_lines='skip',
        )
        before = len(df)
        df = df.dropna(subset=['HB120', 'HH030'])
        df['HB120'] = df['HB120'].astype(int)
        df['HH030'] = df['HH030'].astype(int)
        df = df[(df['HB120'] >= 1) & (df['HH030'] >= 1)]
        for col in ['HB010', 'HB020', 'HB030']:
            df[col] = df[col].astype(str)
        print(f"    H file {country}/{year}: {len(df):,} rows (dropped {before - len(df)} NaN/invalid)")
        return df
    except Exception as exc:
        print(f"  [ERROR] H file {country}/{year}: {exc}")
        return None


def load_region_weights(year: int, country: str) -> pd.DataFrame | None:
    """D file: DB010 DB020 DB030 DB040 DB090"""
    path = _path('D', country, year)
    if not os.path.exists(path):
        print(f"  [WARN] D file not found: {path}")
        return None
    try:
        df = pd.read_csv(
            path,
            usecols=['DB010', 'DB020', 'DB030', 'DB040', 'DB090'],
            on_bad_lines='skip',
        )
        before = len(df)
        df = df.dropna(subset=['DB090', 'DB040'])
        df = df[df['DB090'] > 0]
        df['DB040'] = df['DB040'].astype(str).str.strip()
        for col in ['DB010', 'DB020', 'DB030']:
            df[col] = df[col].astype(str)
        print(f"    D file {country}/{year}: {len(df):,} rows (dropped {before - len(df)} NaN/zero-weight)")
        return df
    except Exception as exc:
        print(f"  [ERROR] D file {country}/{year}: {exc}")
        return None


def compute_decile_thresholds(hh_df: pd.DataFrame) -> dict | None:
    """
    Weighted decile thresholds for equi_disp_inc using DB090.
    Returns dict {1: cut_10pct, ..., 9: cut_90pct}  (9 cut-points → 10 deciles).
    """
    valid = hh_df.dropna(subset=['equi_disp_inc', 'DB090']).copy()
    if len(valid) == 0:
        return None
    valid = valid.sort_values('equi_disp_inc').reset_index(drop=True)
    valid['cum_w']   = valid['DB090'].cumsum()
    total_w          = valid['DB090'].sum()
    valid['cum_pct'] = valid['cum_w'] / total_w

    thresholds = {}
    for d in range(1, N_DECILES):
        target = d / N_DECILES
        idx    = (valid['cum_pct'] - target).abs().idxmin()
        thresholds[d] = valid.loc[idx, 'equi_disp_inc']
    return thresholds


def assign_decile(income, thresholds: dict) -> int | float:
    if pd.isna(income) or thresholds is None:
        return np.nan
    for d in range(1, N_DECILES):
        if income <= thresholds[d]:
            return d
    return N_DECILES


# ============================================================================
# YEAR-LEVEL PIPELINE  (merge H + R + D, compute equiv income & deciles)
# ============================================================================

def load_year_data(year: int, country: str) -> pd.DataFrame | None:
    """
    Returns one row per household with:
      HB120, HH030, DB040, DB090, decile
    Decile thresholds are computed at national level for this year.
    """
    hh    = load_household_full(year, country)
    dw    = load_region_weights(year, country)
    equiv = compute_oecd_equiv_size(year, country)

    if hh is None or dw is None:
        return None

    # --- Equivalized income ---
    if equiv is not None and hh['HY020'].notna().sum() > 0:
        hh = hh.merge(equiv, on=['HB010', 'HB020', 'HB030'], how='left')
        hh['equi_disp_inc'] = hh['HY020'] / hh['equiv_size'].replace(0, np.nan)
    else:
        hh['equi_disp_inc'] = np.nan

    # Temporarily join weights to compute national decile thresholds
    hh_for_deciles = hh.merge(
        dw[['DB010', 'DB020', 'DB030', 'DB090']],
        left_on=['HB010', 'HB020', 'HB030'],
        right_on=['DB010', 'DB020', 'DB030'],
        how='left',
    )
    thresholds = compute_decile_thresholds(hh_for_deciles)
    hh['decile'] = hh['equi_disp_inc'].apply(lambda x: assign_decile(x, thresholds))

    # --- Attach region & final weight ---
    merged = hh.merge(
        dw[['DB010', 'DB020', 'DB030', 'DB040', 'DB090']],
        left_on=['HB010', 'HB020', 'HB030'],
        right_on=['DB010', 'DB020', 'DB030'],
        how='inner',
    )
    merged['year'] = year

    n_total  = len(merged)
    n_decile = merged['decile'].notna().sum()
    print(f"    {country}/{year}: {n_total:,} HH  |  {n_decile:,} with decile  "
          f"|  {merged['DB040'].nunique()} NUTS2 regions")

    return merged[['year', 'HB120', 'HH030', 'DB040', 'DB090', 'decile']]


# ============================================================================
# POOLING
# ============================================================================

def load_pooled_data(years: list[int], country: str) -> pd.DataFrame | None:
    all_dfs = []
    for year in years:
        print(f"\n  Year {year}:")
        df = load_year_data(year, country)
        if df is not None:
            all_dfs.append(df)

    if not all_dfs:
        return None

    combined = pd.concat(all_dfs, ignore_index=True)
    print(f"\n  Pooled total : {len(combined):,} household-year observations")
    print(f"  NUTS2 regions: {sorted(combined['DB040'].unique())}")
    print(f"  Decile coverage: {combined['decile'].notna().mean() * 100:.1f}%")
    return combined


# ============================================================================
# CAPPING & MATRIX COMPUTATION
# ============================================================================

def apply_caps(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df['HB120'] = df['HB120'].clip(upper=MAX_HH_SIZE)
    df['HH030'] = df['HH030'].clip(upper=MAX_ROOMS)
    return df


def axis_label(value: int, max_val: int) -> str:
    return f'{value}+' if value == max_val else str(value)


def _compute_matrix_pair(grp: pd.DataFrame, all_sizes, all_rooms) -> dict:
    """Given a subset (region × decile), compute joint + conditional DataFrames."""
    total_w     = grp['DB090'].sum()
    joint       = pd.DataFrame(0.0, index=all_sizes, columns=all_rooms)
    conditional = pd.DataFrame(0.0, index=all_sizes, columns=all_rooms)

    for hh_size, sub_size in grp.groupby('HB120'):
        if hh_size not in all_sizes:
            continue
        size_w = sub_size['DB090'].sum()
        for rooms, sub_cell in sub_size.groupby('HH030'):
            if rooms not in all_rooms:
                continue
            cell_w = sub_cell['DB090'].sum()
            if total_w > 0:
                joint.loc[hh_size, rooms] = cell_w / total_w
            if size_w > 0:
                conditional.loc[hh_size, rooms] = cell_w / size_w

    joint.index.name         = 'household_size'
    joint.columns.name       = 'n_rooms'
    conditional.index.name   = 'household_size'
    conditional.columns.name = 'n_rooms'

    return {
        'joint':        joint,
        'conditional':  conditional,
        'n_obs':        len(grp),
        'total_weight': total_w,
    }


def compute_probability_matrices(df: pd.DataFrame) -> dict:
    """
    Returns nested dict:
      results[level][decile_label] = {'joint', 'conditional', 'n_obs', 'total_weight'}
    level is NATIONAL_LABEL or a NUTS2 code.
    decile_label is 'ALL' or 'D01'..'D10'.
    """
    df = apply_caps(df)

    all_sizes = sorted(df['HB120'].unique())
    all_rooms = sorted(df['HH030'].unique())

    results: dict = {}

    # ---- 1. National level ----
    results[NATIONAL_LABEL] = {}
    results[NATIONAL_LABEL]['ALL'] = _compute_matrix_pair(df, all_sizes, all_rooms)
    for decile in range(1, N_DECILES + 1):
        grp = df[df['decile'] == decile]
        if len(grp) == 0:
            continue
        results[NATIONAL_LABEL][f'D{decile:02d}'] = _compute_matrix_pair(grp, all_sizes, all_rooms)
    print(f"  National: {len(results[NATIONAL_LABEL])} matrices (ALL + {N_DECILES} deciles)")

    # ---- 2. NUTS2 level ----
    for region, region_grp in df.groupby('DB040'):
        results[region] = {}
        results[region]['ALL'] = _compute_matrix_pair(region_grp, all_sizes, all_rooms)
        for decile in range(1, N_DECILES + 1):
            grp = region_grp[region_grp['decile'] == decile]
            if len(grp) < 10:   # skip cells with too few observations
                continue
            results[region][f'D{decile:02d}'] = _compute_matrix_pair(grp, all_sizes, all_rooms)

    n_nuts2 = len(results) - 1
    print(f"  NUTS2   : {n_nuts2} regions processed")
    return results


# ============================================================================
# HOUSEHOLD SIZE DISTRIBUTION
# ============================================================================

def compute_hh_size_distribution(df: pd.DataFrame) -> pd.DataFrame:
    """
    Weighted share of each household size (after capping) per NUTS2 region
    and at national level.

    Returns a DataFrame with columns:
      region | level_type | household_size | share
    where share sums to 1.0 within each region.
    """
    df = apply_caps(df)
    all_sizes = sorted(df['HB120'].unique())

    rows = []

    def _dist(grp, region, level_type):
        total_w = grp['DB090'].sum()
        if total_w == 0:
            return
        for sz in all_sizes:
            sz_w = grp.loc[grp['HB120'] == sz, 'DB090'].sum()
            rows.append({
                'region':         region,
                'level_type':     level_type,
                'household_size': axis_label(sz, MAX_HH_SIZE),
                'share':          round(float(sz_w / total_w), 6),
                'n_obs':          int((grp['HB120'] == sz).sum()),
            })

    # National
    _dist(df, NATIONAL_LABEL, 'national')

    # Per NUTS2
    for region, grp in df.groupby('DB040'):
        _dist(grp, region, 'nuts2')

    return pd.DataFrame(rows)


# ============================================================================
# EXPORT — Long CSV
# ============================================================================

def export_hh_size_distribution(
    dist_df: pd.DataFrame, years: list[int], out_dir: str, prefix: str
) -> None:
    """Export household size distribution as CSV + one Excel sheet per region."""
    years_str = '-'.join(str(y) for y in years)
    dist_df['years_pooled'] = years_str
    dist_df['country']      = COUNTRY

    # --- CSV ---
    csv_path = os.path.join(out_dir, f'{prefix}_hh_size_distribution.csv')
    dist_df.to_csv(csv_path, index=False)
    print(f"  CSV saved:  {os.path.basename(csv_path)}")

    # --- Excel: wide pivot (regions as rows, sizes as columns) ---
    pivot = dist_df.pivot_table(
        index=['level_type', 'region'],
        columns='household_size',
        values='share',
        aggfunc='first',
    ).reset_index()
    pivot.columns.name = None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'HH size distribution'
    ws.cell(1, 1, f'Household size distribution — France — pooled {min(years)}–{max(years)}'
            ).font = Font(bold=True, size=11)
    ws.cell(2, 1, 'Share sums to 1.0 within each row (weighted by DB090)'
            ).font = Font(italic=True, size=9)

    headers = list(pivot.columns)
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(4, col_idx, h)
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center')

    for row_idx, row in enumerate(pivot.itertuples(index=False), start=5):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row_idx, col_idx, val)

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 2, 12)

    xlsx_path = os.path.join(out_dir, f'{prefix}_hh_size_distribution.xlsx')
    wb.save(xlsx_path)
    print(f"  Excel saved: {os.path.basename(xlsx_path)}")


def export_long_csv(results: dict, years: list[int], out_dir: str, prefix: str) -> None:
    rows = []
    years_str = '-'.join(str(y) for y in years)
    for level, deciles in results.items():
        level_type = 'national' if level == NATIONAL_LABEL else 'nuts2'
        for decile_label, mats in deciles.items():
            for prob_type in ('joint', 'conditional'):
                mat = mats[prob_type]
                for hh_size in mat.index:
                    for rooms in mat.columns:
                        rows.append({
                            'country':          COUNTRY,
                            'level_type':       level_type,
                            'region':           level,
                            'years_pooled':     years_str,
                            'decile':           decile_label,
                            'household_size':   hh_size,
                            'n_rooms':          rooms,
                            'probability_type': prob_type,
                            'probability':      round(float(mat.loc[hh_size, rooms]), 6),
                        })

    out_path = os.path.join(out_dir, f'{prefix}_all_levels_deciles.csv')
    pd.DataFrame(rows).to_csv(out_path, index=False)
    print(f"  CSV saved:  {os.path.basename(out_path)}")


# ============================================================================
# EXPORT — Excel helpers + two workbooks
# ============================================================================

_HEADER_FILL = PatternFill('solid', fgColor='1F497D')
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_BOLD_FONT   = Font(bold=True)


def _write_matrix_sheet(
    wb: openpyxl.Workbook,
    sheet_name: str,
    mat: pd.DataFrame,
    title: str,
    n_obs: int,
) -> None:
    ws = wb.create_sheet(title=sheet_name[:31])
    ws.cell(1, 1, title).font       = Font(bold=True, size=11)
    ws.cell(2, 1, f'n HH (unweighted): {n_obs:,}').font = Font(italic=True, size=9)

    ws.cell(4, 1, 'HH_size \\ Rooms').font = _BOLD_FONT
    for col_idx, rooms in enumerate(mat.columns, start=2):
        cell = ws.cell(4, col_idx, axis_label(rooms, mat.columns.max()))
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center')

    vmax = float(mat.values.max()) if mat.values.max() > 0 else 1.0
    for row_idx, hh_size in enumerate(mat.index, start=5):
        ws.cell(row_idx, 1, axis_label(hh_size, mat.index.max())).font = _BOLD_FONT
        for col_idx, rooms in enumerate(mat.columns, start=2):
            val  = float(mat.loc[hh_size, rooms])
            cell = ws.cell(row_idx, col_idx, round(val, 4))
            cell.alignment = Alignment(horizontal='center')
            intensity = int(255 - (val / vmax) * 200) if vmax > 0 else 255
            intensity = max(0, min(255, intensity))
            cell.fill = PatternFill('solid', fgColor=f'FF{intensity:02X}{intensity:02X}')

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 2, 10)


def export_excel_national(results: dict, years: list[int], out_dir: str, prefix: str) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    national = results.get(NATIONAL_LABEL, {})
    if not national:
        return
    for label in (['ALL'] + [f'D{d:02d}' for d in range(1, N_DECILES + 1)]):
        if label not in national:
            continue
        mats       = national[label]
        decile_str = 'All deciles' if label == 'ALL' else f'Income decile {label}'
        _write_matrix_sheet(
            wb, sheet_name=label, mat=mats['conditional'],
            title=f'P(rooms | hh_size) — France national — {decile_str}  [{min(years)}–{max(years)}]',
            n_obs=mats['n_obs'],
        )
    out_path = os.path.join(out_dir, f'{prefix}_national_matrices.xlsx')
    wb.save(out_path)
    print(f"  Excel (national) saved: {os.path.basename(out_path)}")


def export_excel_nuts2(results: dict, years: list[int], out_dir: str, prefix: str) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_sum = wb.create_sheet(title='Summary')
    ws_sum.append(['NUTS2', 'Decile', 'N obs', 'Total weight', 'Years pooled'])
    years_str = '-'.join(str(y) for y in years)

    regions = sorted(k for k in results if k != NATIONAL_LABEL)
    for region in regions:
        for label, mats in results[region].items():
            ws_sum.append([region, label, mats['n_obs'],
                           round(mats['total_weight'], 0), years_str])

    for region in regions:
        for label in (['ALL'] + [f'D{d:02d}' for d in range(1, N_DECILES + 1)]):
            if label not in results[region]:
                continue
            mats       = results[region][label]
            decile_str = 'All deciles' if label == 'ALL' else f'Decile {label}'
            _write_matrix_sheet(
                wb, sheet_name=f'{region}_{label}', mat=mats['conditional'],
                title=(f'P(rooms | hh_size) — NUTS2: {region} — {decile_str}  '
                       f'[{COUNTRY}, {min(years)}–{max(years)}]'),
                n_obs=mats['n_obs'],
            )

    out_path = os.path.join(out_dir, f'{prefix}_nuts2_matrices.xlsx')
    wb.save(out_path)
    print(f"  Excel (NUTS2)    saved: {os.path.basename(out_path)}")


# ============================================================================
# EXPORT — Heatmaps
# ============================================================================

def _heatmap(ax, mat: pd.DataFrame, vmax: float) -> object:
    data = mat.values.astype(float)
    im   = ax.imshow(data, aspect='auto', cmap='YlOrRd', vmin=0, vmax=max(vmax, 1e-6))
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            val        = data[i, j]
            text_color = 'white' if vmax > 0 and val > vmax * 0.65 else 'black'
            ax.text(j, i, f'{val * 100:.1f}', ha='center', va='center',
                    fontsize=7, color=text_color)
    return im


def export_heatmaps(results: dict, years: list[int], out_dir: str, prefix: str) -> None:
    hm_dir = os.path.join(out_dir, 'heatmaps')
    os.makedirs(hm_dir, exist_ok=True)

    n_saved = 0
    for level, deciles in results.items():
        for decile_label, mats in deciles.items():
            mat = mats['conditional']
            if mat.empty or mats['n_obs'] < 10:
                continue

            n_rows_mat, n_cols_mat = mat.shape
            fig, ax = plt.subplots(
                figsize=(max(7, n_cols_mat * 1.1), max(4, n_rows_mat * 0.9))
            )
            vmax = float(mat.values.max())
            im   = _heatmap(ax, mat, vmax=min(vmax * 1.05, 1.0))

            ax.set_xticks(range(n_cols_mat))
            ax.set_xticklabels(
                [axis_label(r, mat.columns.max()) for r in mat.columns], fontsize=9)
            ax.set_yticks(range(n_rows_mat))
            ax.set_yticklabels(
                [axis_label(h, mat.index.max()) for h in mat.index], fontsize=9)
            ax.set_xlabel('Number of rooms — HH030', fontsize=10)
            ax.set_ylabel('Household size — HB120',  fontsize=10)

            region_str = 'France (national)' if level == NATIONAL_LABEL else f'NUTS2: {level}'
            decile_str = 'All deciles' if decile_label == 'ALL' else f'Income {decile_label}'
            ax.set_title(
                f'P(Rooms | Household Size) — {region_str} — {decile_str}\n'
                f'(pooled {min(years)}–{max(years)}, n={mats["n_obs"]:,})',
                fontsize=11, fontweight='bold',
            )
            cbar = plt.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
            cbar.set_label('Conditional probability', fontsize=9)

            plt.tight_layout()
            fname = f'{prefix}_{level}_{decile_label}.png'
            fig.savefig(os.path.join(hm_dir, fname), dpi=150,
                        bbox_inches='tight', facecolor='white')
            plt.close(fig)
            n_saved += 1

    print(f"  Heatmaps saved: {n_saved} PNGs  →  heatmaps/")


# ============================================================================
# MAIN
# ============================================================================

def main():
    print('=' * 70)
    print('EU-SILC — Household Size × Rooms  |  National + NUTS2 + Decile  |  France')
    print(f'Variables : HB120 (hh size), HH030 (rooms), DB040 (NUTS2), HY020 (income)')
    print(f'Years     : {YEARS}  (pooled)')
    print(f'Caps      : hh size ≤ {MAX_HH_SIZE}+, rooms ≤ {MAX_ROOMS}+')
    print('=' * 70)

    # 1. Load & merge
    df = load_pooled_data(YEARS, COUNTRY)
    if df is None:
        print('\nERROR: no data loaded — check BASE_DATA_PATH and YEARS.')
        return

    # 2. Diagnostics
    print(f'\nHousehold size range : {df["HB120"].min()} – {df["HB120"].max()}')
    print(f'Room count range     : {df["HH030"].min()} – {df["HH030"].max()}')
    print(f'Decile range         : {df["decile"].min()} – {df["decile"].max()}')

    # 3. Compute matrices
    print('\nComputing probability matrices ...')
    results = compute_probability_matrices(df)

    # 4. Household size distribution
    print('\nComputing household size distributions ...')
    dist_df = compute_hh_size_distribution(df.copy())

    # 5. Export
    print('\nExporting outputs ...')
    prefix = f'fr_{min(YEARS)}_{max(YEARS)}'
    export_hh_size_distribution(dist_df, YEARS, OUTPUT_BASE, prefix=prefix)
    export_long_csv(results, YEARS, OUTPUT_BASE, prefix=prefix)
    export_excel_national(results, YEARS, OUTPUT_BASE, prefix=prefix)
    export_excel_nuts2(results, YEARS, OUTPUT_BASE, prefix=prefix)
    export_heatmaps(results, YEARS, OUTPUT_BASE, prefix=prefix)

    print('\n' + '=' * 70)
    print(f'DONE — all outputs in:\n  {OUTPUT_BASE}')
    print('=' * 70)


if __name__ == '__main__':
    main()
